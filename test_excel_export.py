"""Check that local Forecast formulas preserve the channel/cohort results."""
import io
import numpy as np
import openpyxl
import pytest
from excel_export import build_excel
from planning_model import BASE, CHANNELS, Drivers, forecast


@pytest.mark.parametrize('changed', [False, True])
def test_same_sheet_forecast_reconciles(changed):
    channels = CHANNELS.copy(deep=True)
    drivers = BASE
    if changed:
        channels.loc[0, 'Close rate'] = .07
        channels.loc[1, 'Close rate'] = .61
        channels.loc[2, 'Inbound volume'] = 0
        channels.loc[3, 'Lifetime years'] = 1
        drivers = Drivers(fee=.021)
    payload = build_excel(drivers, channels)
    formulas = openpyxl.load_workbook(io.BytesIO(payload), read_only=True)
    cached = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    fs = list(formulas['Forecast'].values)
    vs = list(cached['Forecast'].values)
    model, _ = forecast(drivers, channels=channels, months=60)
    for m in range(60):
        c = m + 6
        letter = openpyxl.utils.get_column_letter(c + 1)
        value = lambda row: vs[row - 1][c]
        assert fs[10][c] == f'={letter}10*{letter}13'
        assert fs[12][c] == f'=IF({letter}10=0,0,SUMPRODUCT({letter}45:{letter}48,{letter}50:{letter}53)/{letter}10)'
        assert fs[15][c] == f'={letter}14+{letter}11-{letter}15'
        assert fs[23][c] == f'={letter}23*{letter}18/12'
        assert fs[24][c] == f'={letter}19+{letter}20-{letter}21+{letter}22-{letter}24'
        opps = [value(r) for r in range(45, 49)]
        rates = [value(r) for r in range(50, 54)]
        weighted = sum(o*r for o, r in zip(opps, rates))/sum(opps) if sum(opps) else 0
        funded = sum(opps)*weighted
        revenue = value(23)*value(18)/12
        ending = value(19)+value(20)-value(21)+value(22)-revenue
        np.testing.assert_allclose(
            [funded, revenue, ending, value(14)+funded-value(15), revenue-value(12)],
            model.iloc[m][['Funded_accounts','Revenue','Ending_AUM','Active_accounts','Revenue_less_marketing']].astype(float),
            rtol=1e-10, atol=1e-6,
        )
    formulas.close()
    cached.close()
